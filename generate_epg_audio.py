"""Create a Mandarin study track from the EPG interview-question DOCX."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def main(input_path: str, output_path: str) -> None:
    document = Document(input_path)
    if not document.tables:
        raise SystemExit("No question table found in the document.")

    rows = document.tables[0].rows
    if len(rows) < 2:
        raise SystemExit("The question table has no content rows.")

    parts = [
        "EPG访谈题库，学习音频。",
        "每题依次播放问题、关键词和完整回答。请在听到关键词后，先自己组织答案，再核对完整回答。",
        "现在开始。",
    ]
    for index, row in enumerate(rows[1:], start=1):
        question, keywords, answer = (clean(cell.text) for cell in row.cells[:3])
        parts.extend(
            [
                f"第{index}题。问题：{question}",
                f"关键词：{keywords}。",
                "请思考。",
                f"参考回答：{answer}",
            ]
        )
    parts.extend(["本轮学习结束。建议重新播放，重点练习刚才不熟悉的问题。"])
    Path(output_path).write_text("\n\n".join(parts), encoding="utf-8")
    print(f"Wrote {len(rows) - 1} questions to {output_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: generate_epg_audio.py INPUT.docx OUTPUT.txt")
    main(sys.argv[1], sys.argv[2])
