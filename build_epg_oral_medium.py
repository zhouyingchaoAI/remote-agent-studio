import json
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
APP = ROOT / '.worktrees/epg-interview-trainer/interview-trainer'
OUT_AUDIO = ROOT / 'output/audio'
OUT_DOCX = ROOT / 'output/docx/EPG访谈口语化背诵版_中间版.docx'
NODE = Path('/Users/zhouyingchao/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node')

SHORT_QUESTIONS = {
    'epg-01': '你在过程改进中负责什么？', 'epg-02': '公司给你哪些支持？',
    'epg-03': '谁检查过程是否遵循、是否有效？', 'epg-04': '你怎样完善公司的过程和资产？',
    'epg-05': 'EPG小组由哪些人组成？', 'epg-06': '怎样找出过程的强项和弱项？',
    'epg-07': '怎样发现后续改进机会？', 'epg-08': '确定机会后，怎样策划改进？',
    'epg-09': '改进目标怎样关联业务目标？', 'epg-10': '哪些关键过程最能支撑业务目标？',
    'epg-11': '怎样发现新方法和新工具？', 'epg-12': '怎样帮助员工参与过程改进？',
    'epg-13': '怎样把新过程部署到项目？', 'epg-14': '怎样判断改进有没有效果？',
    'epg-15': '模板和工具怎么选？', 'epg-16': '过程资产怎么决定复用、开发还是采购？',
    'epg-17': '公司资产库怎样建设和维护？', 'epg-18': '公司的研发过程体系分几层？',
    'epg-19': '过程和资产是怎么开发出来的？', 'epg-20': '怎样沟通和部署过程资产？',
    'epg-21': '怎样保证过程和组织目标一致？', 'epg-22': '怎样让员工真正使用过程资产？',
    'epg-23': '裁剪指南怎么制定，依据是什么？', 'epg-24': '公司的标准工作环境是什么？',
    'epg-25': '项目主要度量哪些目标？', 'epg-26': '项目度量信息怎么确定？',
    'epg-27': '度量数据怎么收集？', 'epg-28': '度量数据怎么分析？',
    'epg-29': '度量数据怎么存储？', 'epg-30': '目标有偏差时怎么纠正？',
    'epg-31': '怎样保证度量数据可靠？', 'epg-32': '项目数据怎样汇总到组织库？',
    'epg-33': '度量结果怎样传达给全员？',
}

# 一句口述版：保留角色、依据、动作或产物，但删除扩展解释和重复修饰。
MEDIUM_ANSWERS = {
    'epg-01': '我是EPG组长，负责组织和推动过程改进、维护标准过程与资产，确保改进支撑业务目标。',
    'epg-02': '公司提供人员、软硬件和模板指南等过程资产，并安排质量体系、CMMI、OSSP和工具等培训。',
    'epg-03': 'QA检查项目对标准过程的遵循情况，高层监督过程改进工作的开展和实际效果。',
    'epg-04': '我结合项目反馈、QA报告、度量数据和高层战略，持续更新公司的标准过程和过程资产。',
    'epg-05': 'EPG依据《过程改进计划》组建，由高层、PM、CM、QA、OT等跨角色成员组成，并明确职责。',
    'epg-06': '我们开展内部过程诊断，必要时邀请外部咨询，通过差距分析识别过程强项、弱项和改进重点。',
    'epg-07': '我们从差距分析问题、QA报告、项目反馈和度量数据中，结合研发业务目标识别改进机会。',
    'epg-08': '我们在《过程改进计划》中明确改进目标、试点和部署计划、责任人，以及风险和应对措施。',
    'epg-09': '过程改进目标都对应业务目标，并用缺陷率、成本偏差和客户满意度等量化指标验证。',
    'epg-10': '成本和工作量目标重点看计划监控，质量和效率目标重点看QA、同行评审和测试过程。',
    'epg-11': '我们通过头脑风暴、技术讨论会和EPG例会探索新方法工具，先在项目试点验证再决定推广。',
    'epg-12': '我参与OSSP文件编写和培训，并在项目中提供指导、监督执行和收集反馈，促进员工参与。',
    'epg-13': '新过程先选择合适项目试点，收集反馈后修改完善，再正式推广，并持续协调资源和跟踪问题。',
    'epg-14': '我们从度量指标是否达标、QA报告反映的遵循情况和过程资产使用情况三方面评估改进效果。',
    'epg-15': '我们优先从组织资产库选用模板和工具，再结合客户要求和项目实际进行裁剪或补充。',
    'epg-16': '过程资产优先复用；无法复用时，综合成本、进度、周期、风险和技术成熟度决定开发或采购。',
    'epg-17': '资产库以降低开发成本、提高质量和规范过程为目标，并融合CMMI、ISO等体系要求持续维护。',
    'epg-18': '研发过程体系分三层：顶层是方针和业务目标，中层是过程及模板指南，底层是组织资产库。',
    'epg-19': '我们参考既有体系，结合质量体系要求、客户需求、公司战略和行业最佳实践开发过程和资产。',
    'epg-20': '我通过EPG例会和项目立项会沟通部署并提供指导，项目结束后将项目资产汇总到组织资产库。',
    'epg-21': '我们结合组织目标和项目数据制定标准过程，推广前先在项目试点，验证过程与目标的一致性。',
    'epg-22': '我们持续更新并记录过程资产修订历史，开展培训，同时通过GitLab等工具为员工开放使用权限。',
    'epg-23': 'EPG共同制定《裁剪指南》，按行业标准、CMMI要求、客户类型和项目规模等依据裁剪标准过程。',
    'epg-24': '《工作环境标准》统一规定办公环境、软硬件、开发平台和工具，由EPG定期维护更新。',
    'epg-25': '项目度量目标包括控制缺陷率在0.03/KLOC、成本偏差在±10%以内，并提升客户满意度。',
    'epg-26': '我们根据度量目标确定进度、工作量、问题数和缺陷率等度量项，并记录在《度量数据表》中。',
    'epg-27': '《度量计划》明确数据来源、收集时间、责任人和验证方式，例如缺陷来自测试报告、工时来自日报。',
    'epg-28': '我们按《度量计划》进行目标对比、偏差原因和数据分布分析，进度成本用挣值法，工作量用偏差法。',
    'epg-29': '度量数据收集后记录在《度量数据表》中，由CM放入配置库统一管理。',
    'epg-30': '我们将问题和纠正措施记录在《问题管理表》中，协调人员、调整计划或安排培训，并反馈给EPG。',
    'epg-31': 'PM负责抽查数据准确性，QA按检查单检查，EPG在分析前再次复核，三层保证数据质量。',
    'epg-32': '项目把度量结果记录在《度量数据表》并提交组织度量库，由CM和EPG共同维护。',
    'epg-33': '我们在管理层会议分享度量数据、分析结果和改进结论，并上传组织度量库供研发人员查阅。',
}


def extract_questions():
    script = "import('./src/questionBank.js').then(m => process.stdout.write(JSON.stringify(m.QUESTIONS)))"
    result = subprocess.run([str(NODE), '-e', script], cwd=APP, check=True, capture_output=True, text=True)
    questions = json.loads(result.stdout)
    if len(questions) != 33 or set(q['id'] for q in questions) != set(MEDIUM_ANSWERS):
        raise ValueError('题库或中间版答案映射不完整。')
    return questions


def set_font(run, size=9, bold=False, color=None):
    run.font.name = 'PingFang SC'
    for key in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        run._element.rPr.rFonts.set(qn(key), 'PingFang SC')
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement('w:tcMar')
    for edge, value in [('top', '80'), ('bottom', '80'), ('start', '120'), ('end', '120')]:
        part = OxmlElement(f'w:{edge}')
        part.set(qn('w:w'), value)
        part.set(qn('w:type'), 'dxa')
        node.append(part)
    tc_pr.append(node)


def build_docx(questions):
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    for attr in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
        setattr(section, attr, Inches(1))
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(3)
    set_font(title.add_run('EPG 访谈口语化背诵版（中间版）'), 19, True, (31, 77, 120))
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.space_after = Pt(12)
    set_font(intro.add_run('一题一整句：顺口作答，同时保留关键依据和动作'), 10, False, (75, 85, 99))

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    widths = [2880, 2160, 4320]
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn('w:w'), str(width))
    for cell, label in zip(table.rows[0].cells, ['简化问题', '关键词', '中间版口述答案']):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EEF5'); cell._tc.get_or_add_tcPr().append(shade)
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.25
        set_font(p.add_run(label), 9, True, (11, 37, 69))
    for index, item in enumerate(questions, start=1):
        row = table.add_row()
        values = [f'{index}. {SHORT_QUESTIONS[item["id"]]}', item['memoryCue'], MEDIUM_ANSWERS[item['id']]]
        for cell, width, value in zip(row.cells, widths, values):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_w = OxmlElement('w:tcW'); tc_w.set(qn('w:w'), str(width)); tc_w.set(qn('w:type'), 'dxa'); cell._tc.get_or_add_tcPr().append(tc_w)
            p = cell.paragraphs[0]; p.paragraph_format.line_spacing = 1.25
            set_font(p.add_run(value), 9)
    doc.save(OUT_DOCX)


def build_audio(questions):
    OUT_AUDIO.mkdir(parents=True, exist_ok=True)
    text = ['EPG访谈口语化背诵中间版。每题依次朗读：简化问题、关键词和口述答案。']
    markdown = ['# EPG 访谈口语化背诵版（中间版）', '']
    for index, item in enumerate(questions, start=1):
        answer = MEDIUM_ANSWERS[item['id']]
        question = SHORT_QUESTIONS[item['id']]
        cue = item['memoryCue'].replace('｜', '，')
        text += [f'第 {index} 题。', f'问题：{question}', f'关键词：{cue}', f'回答：{answer}', '。']
        markdown += [f'## 第 {index} 题', f'- 简化问题：{question}', f'- 关键词：{item["memoryCue"]}', f'- 中间版口述答案：{answer}', '']
    txt = OUT_AUDIO / 'EPG访谈口语化背诵稿_中间版.txt'
    txt.write_text('\n'.join(text), encoding='utf-8')
    (OUT_AUDIO / 'EPG访谈口语化背诵稿_中间版.md').write_text('\n'.join(markdown), encoding='utf-8')
    aiff = OUT_AUDIO / 'EPG访谈口语化背诵版_中间版.aiff'
    mp3 = OUT_AUDIO / 'EPG访谈口语化背诵版_中间版.mp3'
    subprocess.run(['say', '-v', 'Tingting', '-r', '180', '-f', str(txt), '-o', str(aiff)], check=True)
    subprocess.run(['ffmpeg', '-y', '-loglevel', 'error', '-i', str(aiff), '-codec:a', 'libmp3lame', '-b:a', '128k', str(mp3)], check=True)
    aiff.unlink(missing_ok=True)


def main():
    questions = extract_questions()
    build_docx(questions)
    build_audio(questions)
    print(OUT_DOCX)
    print(OUT_AUDIO / 'EPG访谈口语化背诵版_中间版.mp3')


if __name__ == '__main__':
    main()
