import json
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
APP = ROOT / '.worktrees/epg-interview-trainer/interview-trainer'
OUT = ROOT / 'output/docx/EPG访谈口语化背诵版.docx'
NODE = Path('/Users/zhouyingchao/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node')

SHORT_QUESTIONS = {
    'epg-01': '你在过程改进中负责什么？', 'epg-02': '公司给你哪些支持？',
    'epg-03': '谁检查过程是否遵循、是否有效？', 'epg-04': '你怎样完善公司的过程和资产？',
    'epg-05': 'EPG小组由哪些人组成？', 'epg-06': '怎样找出过程的强项和弱项？',
    'epg-07': '怎样发现后续改进机会？', 'epg-08': '确定机会后，怎样策划改进？',
    'epg-09': '改进目标怎样关联业务目标？', 'epg-10': '哪些关键过程最能支撑业务目标？',
    'epg-11': '怎样发现新方法和新工具？', 'epg-12': '怎样帮助员工参与过程改进？',
    'epg-13': '怎样把新过程部署到项目？', 'epg-14': '怎样判断改进有没有效果？',
    'epg-15': '模板和工具怎么选？', 'epg-16': '过程资产怎么决定复用、开发还是采购？',
    'epg-17': '公司资产库怎样建设和维护？', 'epg-18': '公司的研发过程体系分几层？',
    'epg-19': '过程和资产是怎么开发出来的？', 'epg-20': '怎样沟通和部署过程资产？',
    'epg-21': '怎样保证过程和组织目标一致？', 'epg-22': '怎样让员工真正使用过程资产？',
    'epg-23': '裁剪指南怎么制定，依据是什么？', 'epg-24': '公司的标准工作环境是什么？',
    'epg-25': '项目主要度量哪些目标？', 'epg-26': '项目度量信息怎么确定？',
    'epg-27': '度量数据怎么收集？', 'epg-28': '度量数据怎么分析？',
    'epg-29': '度量数据怎么存储？', 'epg-30': '目标有偏差时怎么纠正？',
    'epg-31': '怎样保证度量数据可靠？', 'epg-32': '项目数据怎样汇总到组织库？',
    'epg-33': '度量结果怎样传达给全员？',
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement('w:shd')
    shade.set(qn('w:fill'), fill)
    tc_pr.append(shade)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn('w:tcW'))
    if tc_w is None:
        tc_w = OxmlElement('w:tcW')
        tc_pr.append(tc_w)
    tc_w.set(qn('w:w'), str(width_dxa))
    tc_w.set(qn('w:type'), 'dxa')


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in('w:tblW')
    tbl_w.set(qn('w:w'), '9360')
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_layout = OxmlElement('w:tblLayout')
    tbl_layout.set(qn('w:type'), 'fixed')
    tbl_pr.append(tbl_layout)
    indent = OxmlElement('w:tblInd')
    indent.set(qn('w:w'), '120')
    indent.set(qn('w:type'), 'dxa')
    tbl_pr.append(indent)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn('w:w'), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement('w:tcMar')
    for edge, value in [('top', '80'), ('bottom', '80'), ('start', '120'), ('end', '120')]:
        node = OxmlElement(f'w:{edge}')
        node.set(qn('w:w'), value)
        node.set(qn('w:type'), 'dxa')
        margins.append(node)
    tc_pr.append(margins)


def set_font(run, size=9, bold=False, color=None):
    run.font.name = 'Heiti SC'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Heiti SC')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Heiti SC')
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Heiti SC')
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_para(paragraph, before=0, after=0, line=1.25, alignment=None):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if alignment is not None:
        paragraph.alignment = alignment


def extract_questions():
    script = "import('./src/questionBank.js').then(m => process.stdout.write(JSON.stringify(m.QUESTIONS)))"
    result = subprocess.run([str(NODE), '-e', script], cwd=APP, check=True, capture_output=True, text=True)
    questions = json.loads(result.stdout)
    if len(questions) != 33 or set(q['id'] for q in questions) != set(SHORT_QUESTIONS):
        raise ValueError('题库或口语化问题映射不完整。')
    return questions


def build():
    questions = extract_questions()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles['Normal']
    normal.font.name = 'Heiti SC'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Heiti SC')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Heiti SC')
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Heiti SC')
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    set_para(title, after=3, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(title.add_run('EPG 访谈口语化背诵版'), size=20, bold=True, color=(31, 77, 120))
    subtitle = doc.add_paragraph()
    set_para(subtitle, after=12, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(subtitle.add_run('练法：读简化问题 → 看关键词 → 口述短答'), size=10, color=(75, 85, 99))

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    widths = [2880, 2160, 4320]
    headers = ['简化问题', '关键词', '建议短答']
    for cell, label in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, 'E8EEF5')
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        set_para(p, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        set_font(p.add_run(label), size=9, bold=True, color=(11, 37, 69))

    for number, item in enumerate(questions, start=1):
        cells = table.add_row().cells
        values = [f'{number}. {SHORT_QUESTIONS[item["id"]]}', item['memoryCue'].replace('｜', '｜'), item['shortAnswer']]
        for cell, value in zip(cells, values):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            set_para(p)
            set_font(p.add_run(value), size=9)
    set_table_geometry(table, widths)
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build()
