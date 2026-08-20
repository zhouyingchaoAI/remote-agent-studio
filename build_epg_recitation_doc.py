from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = 'EPG访谈题库_背诵版.docx'
items = [
('您在过程改进中的角色和责任是什么？', '角色｜工作｜目的', '我是EPG组长，负责改进过程，支撑业务目标。'),
('公司为您工作的开展提供了哪些资源、培训以及相关过程资产？', '资源｜资产｜培训', '公司提供人、软硬件、过程资产和相关培训。'),
('谁来评价对于标准过程的遵循度和过程的有效性？', 'QA｜高层', 'QA检查过程遵循，高层监督改进效果。'),
('您如何为完善组织的资产和过程做出贡献？', '反馈｜数据｜更新', '我根据反馈和数据，持续更新过程与资产。'),
('过程改进小组的组成结构是怎样的？', '高层｜PM｜CM/QA/OT', 'EPG由高层、PM、CM、QA、OT等组成，职责明确。'),
('您如何评估当前的过程并识别强弱项？', '诊断｜咨询｜差距', '通过内部诊断和差距分析，识别过程强弱项。'),
('您如何识别未来的改进机会？（即：通过什么途径和方式）', '差距｜QA｜反馈/数据', '从差距、QA报告、项目反馈和数据中识别机会。'),
('基于所选取的过程改进机会，您如何来策划未来的改进？', '计划｜试点｜风险', '在改进计划中明确试点、部署、责任和风险。'),
('您如何制定可追溯到业务目标的过程改进目标？', '质量｜成本｜满意度', '改进目标对应质量、成本和客户满意度等业务目标。'),
('研发过程中，有哪些重要过程，最有助于业务目标的实现？', '监控｜QA/评审｜测试', '成本看计划监控，质量和效率看QA、评审、测试。'),
('在改进过程的实施中，有使用哪些新的技术，方法或工具来发现改进机会或识别新的过程？', '讨论｜探索｜试点', '通过讨论会探索新方法，先试点验证再推广。'),
('您提供什么支持和帮助，使大家能够积极地参与过程改进？', 'OSSP｜培训｜指导', '我编写并培训OSSP，在项目中持续指导和收反馈。'),
('您执行了哪些活动来部署过程？', '试点｜反馈｜推广', '过程先试点、收反馈、做修改，再正式推广。'),
('您如何评估已经实施的改进项对实现业务目标的影响？', '指标｜QA｜使用率', '看度量是否达标、QA报告和资产使用情况。'),
('您如何确定研发过程中需要用到哪些模板和工具？', '资产库｜客户｜裁剪', '优先使用资产库，再按客户需求进行裁剪。'),
('您如何确定需要购买，复用或开发哪些过程资产？除了复用之外，有没有购买或者开发一些过程资产？', '复用优先｜成本｜风险', '优先复用，再看成本、进度、风险决定开发或采购。'),
('您使用什么策略来制定和维护公司的资产库？', '降成本｜提质量｜规范', '资产库以降成本、提质量、规范过程为目标持续维护。'),
('请谈一下公司的研发过程体系结构是怎样的？', '方针｜过程｜资产库', '体系分三层：顶层方针、中层过程、底层资产库。'),
('您如何开发支持定义的体系结构所需的资产和过程？', '体系｜客户｜最佳实践', '结合现有体系、客户需求和最佳实践开发资产过程。'),
('您如何与相关人员沟通和部署资产和过程？', 'EPG会｜立项会｜入库', '通过EPG会和立项会沟通，项目结束后资产入库。'),
('如何确保过程与组织目标保持一致？', '目标｜数据｜试点', '结合组织目标和项目数据制定过程，并通过试点验证。'),
('您如何确保公司员工会使用公司的过程资产？', '维护｜培训｜权限', '持续维护资产、开展培训，并通过GitLab开放权限。'),
('请解释如何制定裁剪指南，以及指南的内容', 'EPG｜标准｜客户/规模', 'EPG制定裁剪指南，按标准、客户和规模裁剪。'),
('公司定义的标准工作环境是什么？', '办公｜软硬件｜工具', '工作环境标准统一规定办公、软硬件和开发工具。'),
('您为项目设定的度量目标是什么？', '质量｜成本｜满意度', '度量目标是控制缺陷、成本偏差并提升满意度。'),
('您如何定义项目度量所需的必要信息？', '目标｜度量项｜数据表', '由度量目标确定进度、工作量、缺陷等度量项。'),
('如何收集指定的度量数据？', '来源｜时间｜责任人', '度量计划规定数据来源、时间、责任人和验证方式。'),
('如何分析度量数据？', '目标对比｜原因｜分布', '通过目标对比、原因和分布分析数据。'),
('如何存储度量数据？', '数据表｜CM｜配置库', '数据记录在度量数据表，由CM放入配置库。'),
('如果实际情况与目标实现产生偏差，该采取什么纠正措施？', '问题表｜纠正｜EPG', '将问题和纠正措施记入问题表，并反馈给EPG。'),
('您如何确保度量数据的质量？', 'PM｜QA｜EPG', 'PM抽查、QA检查、EPG复核，三层保障数据质量。'),
('您如何在组织库中汇总项目的度量数据？', '数据表｜组织库｜CM+EPG', '项目提交度量数据表，由CM和EPG维护组织度量库。'),
('您如何将度量结果传达给组织中的所有人员？', '管理层会｜结论｜度量库', '在管理层会议分享结论，并上传组织度量库。'),
]

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd'); el.set(qn('w:fill'), fill); pr.append(el)

def margins(cell):
    pr = cell._tc.get_or_add_tcPr(); mar = OxmlElement('w:tcMar')
    for k,v in [('top',100),('bottom',100),('start',140),('end',140)]:
        e=OxmlElement(f'w:{k}'); e.set(qn('w:w'),str(v)); e.set(qn('w:type'),'dxa'); mar.append(e)
    pr.append(mar)

def put(cell, text, bold=False, color='000000', size=10):
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.15
    r=p.add_run(text); r.bold=bold; r.font.name='Hiragino Sans GB'; r.font.size=Pt(size)
    for attr in ('ascii','hAnsi','eastAsia'): r._element.rPr.rFonts.set(qn('w:'+attr),'Hiragino Sans GB')
    r.font.color.rgb=RGBColor.from_string(color); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cell)

doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.6); sec.bottom_margin=Inches(.6); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
title=doc.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.paragraph_format.space_after=Pt(3)
r=title.add_run('EPG访谈题库 · 背诵版'); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=RGBColor(31,77,120)
sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; sub.paragraph_format.space_after=Pt(10)
sub.add_run('背法：先看“提示词”，再说“短答”；每题只背一个动作句。').font.size=Pt(10)
table=doc.add_table(rows=1,cols=3); table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.LEFT
widths=[3300,1900,4160]
pr=table._tbl.tblPr; layout=OxmlElement('w:tblLayout'); layout.set(qn('w:type'),'fixed'); pr.append(layout)
grid=table._tbl.tblGrid
for i,w in enumerate(widths): grid.gridCol_lst[i].set(qn('w:w'),str(w))
for i,h in enumerate(['固定问题','提示词（看一眼）','短答（直接说）']): shade(table.rows[0].cells[i],'E8EEF5'); put(table.rows[0].cells[i],h,True,'1F4D78',10)
for q,k,a in items:
    cells=table.add_row().cells
    for i,(text,bold,color) in enumerate([(q,True,'1F4D78'),(k,False,'7A5A00'),(a,False,'000000')]):
        put(cells[i],text,bold,color,9.5); cells[i].width=Inches(widths[i]/1440)
for row in table.rows:
    trPr=row._tr.get_or_add_trPr(); cs=OxmlElement('w:cantSplit'); trPr.append(cs)
headerPr=table.rows[0]._tr.get_or_add_trPr(); hdr=OxmlElement('w:tblHeader'); hdr.set(qn('w:val'),'true'); headerPr.append(hdr)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run('EPG访谈背诵版').font.size=Pt(8)
doc.save(OUT)
print(OUT)
